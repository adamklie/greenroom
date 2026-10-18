"""Import a gig setlist from an .xlsx sheet.

Expected layout (the band leader's set-list sheets):

    SET LIST 101826                         <- title rows above the header
    Gaylord ICC event
    #  | Song | Artist | Notes | Tab | Track | Chart   <- header row ("Song" cell)
    SET 1 7:30 - 9:30                       <- text row with no song = section
    1  | Anyway You Want It | Journey | ... | <hyperlinks>
    Potential EXTRAS:
       | Come On Eileen | ...

Columns are found by header name. Any column whose cells carry hyperlinks is a
link column; other unrecognised text columns are ignored. The Song cell's own
hyperlink is treated as a "listen" link.
"""

from __future__ import annotations

from io import BytesIO
from urllib.parse import urlparse

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Setlist, SetlistItem, Song
from app.routers.dedup import _normalize

# Charts are gig-specific (numbered per event), so they always stay on the item.
ITEM_ONLY_KINDS = {"chart"}


class SetlistImportError(ValueError):
    pass


def _kind_for(header: str, url: str) -> str:
    h = header.lower()
    if "tab" in h:
        return "tab"
    if "track" in h:
        return "track"
    if "chart" in h:
        return "chart"
    host = urlparse(url).netloc.lower()
    if "ultimate-guitar" in host:
        return "tab"
    if "spotify" in host or "youtube" in host or "youtu.be" in host:
        return "listen"
    return "other"


def _text(cell) -> str:
    return str(cell.value).strip() if cell.value is not None else ""


def parse_setlist_xlsx(data: bytes) -> dict:
    """Parse the sheet into {name, description, rows}. Each row is
    {title, artist, notes, section, links: [{kind, label, url}]}."""
    try:
        # Not read_only: read-only mode drops cell hyperlinks.
        ws = load_workbook(BytesIO(data)).worksheets[0]
    except Exception as e:  # openpyxl raises a zoo of types for bad input
        raise SetlistImportError(f"Couldn't read workbook: {e}") from e

    rows = list(ws.iter_rows())
    header_idx = next(
        (i for i, r in enumerate(rows) if any(_text(c).lower() == "song" for c in r)),
        None,
    )
    if header_idx is None:
        raise SetlistImportError("No header row with a 'Song' column found")

    headers = {c.column: _text(c) for c in rows[header_idx] if _text(c)}
    by_name = {h.lower(): col for col, h in headers.items()}
    song_col = by_name["song"]
    num_col, artist_col, notes_col = by_name.get("#"), by_name.get("artist"), by_name.get("notes")
    named = {song_col, num_col, artist_col, notes_col}

    title_lines = [
        _text(c) for r in rows[:header_idx] for c in r if _text(c)
    ]
    out_rows: list[dict] = []
    section: str | None = None
    for r in rows[header_idx + 1:]:
        cells = {c.column: c for c in r}
        song_cell = cells.get(song_col)
        title = _text(song_cell) if song_cell is not None else ""
        if not title:
            # A text row without a song is a section label.
            label = next((_text(c) for c in r if _text(c)), "")
            if label:
                section = label.rstrip(":").strip()
            continue

        links = []
        if song_cell.hyperlink and song_cell.hyperlink.target:
            links.append({"kind": "listen", "label": None, "url": song_cell.hyperlink.target})
        for col, header in headers.items():
            if col in named:
                continue
            c = cells.get(col)
            if c is None or not c.hyperlink or not c.hyperlink.target:
                continue
            url = c.hyperlink.target
            links.append({"kind": _kind_for(header, url), "label": _text(c) or header, "url": url})

        out_rows.append({
            "title": title,
            "artist": _text(cells[artist_col]) or None if artist_col in cells else None,
            "notes": _text(cells[notes_col]) or None if notes_col in cells else None,
            "section": section,
            "links": links,
        })

    if not out_rows:
        raise SetlistImportError("No songs found below the header row")
    return {
        "name": title_lines[0] if title_lines else "Imported setlist",
        "description": " · ".join(title_lines[1:]) or None,
        "rows": out_rows,
    }


def _find_song(db: Session, title: str, artist: str | None) -> Song | None:
    """Match on normalized title+artist; fall back to a unique title match.
    Queries are scoped to the active project by the ORM listener."""
    nt, na = _normalize(title), _normalize(artist or "")
    candidates = [
        s for s in db.query(Song).filter(Song.status != "deleted").all()
        if _normalize(s.title) == nt
    ]
    exact = [s for s in candidates if _normalize(s.artist or "") == na]
    if exact:
        return exact[0]
    return candidates[0] if len(candidates) == 1 else None


def import_setlist(db: Session, data: bytes) -> tuple[Setlist, int, int]:
    parsed = parse_setlist_xlsx(data)
    setlist = Setlist(name=parsed["name"], description=parsed["description"])
    db.add(setlist)
    db.flush()

    matched = created = 0
    for pos, row in enumerate(parsed["rows"]):
        song = _find_song(db, row["title"], row["artist"])
        if song:
            matched += 1
        else:
            song = Song(
                title=row["title"], artist=row["artist"],
                type="cover", status="learning", project="solo",
            )
            db.add(song)
            db.flush()
            created += 1

        song_links = list(song.links or [])
        item_links = []
        for link in row["links"]:
            same_kind = [existing for existing in song_links if existing["kind"] == link["kind"]]
            if link["kind"] in ITEM_ONLY_KINDS:
                item_links.append(link)
            elif not same_kind:
                song_links.append(link)
            elif all(existing["url"] != link["url"] for existing in same_kind):
                item_links.append(link)
        song.links = song_links

        db.add(SetlistItem(
            setlist_id=setlist.id, song_id=song.id, position=pos,
            notes=row["notes"], section=row["section"], links=item_links or None,
        ))

    db.commit()
    db.refresh(setlist)
    return setlist, matched, created
