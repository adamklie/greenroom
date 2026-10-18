"""Setlist sections, links, and .xlsx import."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.models import Song
from tests.test_project_isolation import _as, _h, iso  # noqa: F401  (fixture)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
UG = "https://tabs.ultimate-guitar.com/tab/journey/open-arms-official-1"
UG_OTHER = "https://tabs.ultimate-guitar.com/tab/journey/open-arms-chords-2"


def _link(ws, ref, text, url):
    ws[ref] = text
    ws[ref].hyperlink = url


def _sheet() -> bytes:
    """Same shape as the band leader's sheets: title rows, header, two
    sections, a LEAD column that should be ignored."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "SET LIST 101826"
    ws["A2"] = "Gaylord ICC event"
    for col, h in zip("ABCDEFGH", ["#", "Song", "LEAD", "Artist", "Notes", "UG Official Tab", "TRACK", "CHART"]):
        ws[f"{col}4"] = h
    ws["A5"] = "SET 1 7:30 - 9:30"

    ws["A6"] = 1
    _link(ws, "B6", "Open Arms", "https://open.spotify.com/track/1")
    ws["C6"] = "Desi"
    ws["D6"] = "Journey"
    ws["E6"] = "Elec Intro"
    _link(ws, "F6", "Official Tab", UG_OTHER)
    _link(ws, "G6", "Track", "https://www.dropbox.com/s/a/OPEN%20ARMS%20BKG.m4a?dl=0")
    _link(ws, "H6", "Chart", "https://www.dropbox.com/s/b/GICC1018-06-Open-Arms.pdf?dl=0")

    ws["A7"] = 2
    ws["B7"] = "Jump"
    ws["D7"] = "Van Halen"
    _link(ws, "F7", "Official Tab", "https://tabs.ultimate-guitar.com/tab/van-halen/jump-1")

    ws["A9"] = "Potential EXTRAS:"
    ws["B10"] = "My Sharona"
    ws["D10"] = "The Knack"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, data: bytes, headers=None):
    return client.post(
        "/api/setlists/import",
        files={"file": ("set.xlsx", data, XLSX)},
        headers=headers or {},
    )


def test_import_builds_sections_notes_and_links(client, db):
    # Pre-existing song with a different tab: it should be matched, keep its
    # own tab, and get the sheet's tab as a per-gig override.
    db.add(Song(title="Open Arms", artist="Journey", type="cover", status="rehearsed",
                project="solo", links=[{"kind": "tab", "label": None, "url": UG}]))
    db.commit()

    res = _import(client, _sheet())
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["songs_matched"], body["songs_created"]) == (1, 2)

    sl = body["setlist"]
    assert sl["name"] == "SET LIST 101826"
    assert sl["description"] == "Gaylord ICC event"
    items = sl["items"]
    assert [i["song_title"] for i in items] == ["Open Arms", "Jump", "My Sharona"]
    assert [i["section"] for i in items] == ["SET 1 7:30 - 9:30"] * 2 + ["Potential EXTRAS"]
    assert items[0]["notes"] == "Elec Intro"

    open_arms = items[0]
    assert {link["kind"] for link in open_arms["links"]} == {"tab", "chart"}
    eff = {link["kind"]: link["url"] for link in open_arms["effective_links"]}
    assert eff["tab"] == UG_OTHER  # item override wins
    assert set(eff) == {"tab", "chart", "listen", "track"}

    song = db.query(Song).filter(Song.title == "Open Arms").one()
    db.refresh(song)
    song_kinds = sorted(link["kind"] for link in song.links)
    assert song_kinds == ["listen", "tab", "track"]  # chart stays on the item
    assert [link["url"] for link in song.links if link["kind"] == "tab"] == [UG]

    # LEAD column is ignored entirely.
    assert all("Desi" not in str(i) for i in items)


def test_patch_round_trip_keeps_item_fields(client):
    sl = _import(client, _sheet()).json()["setlist"]
    payload = {"items": [
        {"song_id": i["song_id"], "position": i["position"], "duration_minutes": i["duration_minutes"],
         "notes": i["notes"], "section": i["section"], "links": i["links"]}
        for i in sl["items"]
    ]}
    res = client.patch(f"/api/setlists/{sl['id']}", json=payload)
    assert res.status_code == 200
    for before, after in zip(sl["items"], res.json()["items"]):
        for f in ("notes", "section", "links", "effective_links"):
            assert before[f] == after[f]


def test_setlist_items_carry_song_lyrics(client, db):
    sl = _import(client, _sheet()).json()["setlist"]
    song = db.query(Song).filter(Song.title == "Jump").one()
    song.lyrics = "[Verse]\nI get up"
    db.commit()
    items = client.get(f"/api/setlists/{sl['id']}").json()["items"]
    jump = next(i for i in items if i["song_title"] == "Jump")
    assert jump["song_lyrics"] == "[Verse]\nI get up"
    assert jump["song_type"] == "cover"


def test_import_rejects_sheet_without_header(client):
    wb = Workbook()
    wb.active["A1"] = "nothing here"
    buf = BytesIO()
    wb.save(buf)
    assert _import(client, buf.getvalue()).status_code == 400


def test_import_lands_in_active_project_and_viewer_cannot_import(client, db, iso):  # noqa: F811
    _as(client, iso.ua)
    res = _import(client, _sheet(), _h(iso.pa))
    assert res.status_code == 200, res.text
    db.expunge_all()
    created = db.query(Song).filter(Song.title == "Jump").one()
    assert created.project_id == iso.pa

    _as(client, iso.uc)  # viewer of PA
    assert client.get(f"/api/setlists/{res.json()['setlist']['id']}", headers=_h(iso.pa)).status_code == 200
    assert _import(client, _sheet(), _h(iso.pa)).status_code == 403

    _as(client, iso.ub)  # other project can't see it
    assert client.get(f"/api/setlists/{res.json()['setlist']['id']}", headers=_h(iso.pb)).status_code == 404
